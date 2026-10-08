"""Trich xuat va lam sach text CV (code, khong LLM): docs/ai-service/01 muc 2.2-2.4."""
import io
import re
import unicodedata
import zipfile
from typing import Tuple

from pypdf import PdfReader
from docx import Document

from .api_error import ApiError

MAX_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 10
MAX_ZIP_ENTRIES = 200
MAX_UNZIPPED = 20 * 1024 * 1024
MIN_CHARS = 200
MAX_CHARS = 15_000

_ZERO_WIDTH = re.compile("[​-‏⁠﻿]")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
_URL = re.compile(r"(https?://|www\.)\S+|\b(linkedin|github|facebook)\.com/\S+", re.I)
_PHONE = re.compile(r"(\+84|0)[\s.-]?(3|5|7|8|9)([\s.-]?\d){8}")
_DOB = re.compile(r"((ngày sinh|dob|birth)[^\n\d]{0,15})(\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4})", re.I)
# Mau chi dan nghi la prompt injection (docs 2.4 buoc 5)
INJECTION_PATTERNS = [
    r"ignore (all|previous|above)", r"system prompt", r"you are (now )?", r"\binstruction",
    r"bỏ qua (mọi|các) (hướng dẫn|chỉ dẫn)", r"chấm .{0,20}điểm", r"<\|", r"\[inst\]", r"```",
]
_INJECTION = re.compile("|".join(INJECTION_PATTERNS), re.I)


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ApiError(415, "cv_invalid_format", "PDF co mat khau khong duoc ho tro")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ApiError(413, "cv_too_large", f"PDF qua {MAX_PDF_PAGES} trang")
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except ApiError:
        raise
    except Exception:
        raise ApiError(415, "cv_invalid_format", "Khong doc duoc file PDF")


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNZIPPED:
                raise ApiError(415, "cv_invalid_format", "DOCX vuot gioi han an toan (zip bomb)")
            if "word/document.xml" not in zf.namelist():
                raise ApiError(415, "cv_invalid_format", "Khong phai file DOCX")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except ApiError:
        raise
    except Exception:
        raise ApiError(415, "cv_invalid_format", "Khong doc duoc file DOCX")


def extract_text(data: bytes) -> str:
    """Nhan dien dinh dang theo magic bytes, khong theo duoi file."""
    if len(data) > MAX_BYTES:
        raise ApiError(413, "cv_too_large", "CV vuot 10 MB")
    if data.startswith(b"%PDF-"):
        return _pdf_text(data)
    if data.startswith(b"PK\x03\x04"):
        return _docx_text(data)
    raise ApiError(415, "cv_invalid_format", "Chi chap nhan PDF hoac DOCX")


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = _ZERO_WIDTH.sub("", _CONTROL.sub("", text))
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def redact_pii(text: str) -> str:
    text = _DOB.sub(lambda m: m.group(1) + "[DOB]", text)
    text = _EMAIL.sub("[EMAIL]", text)
    text = _URL.sub("[URL]", text)
    return _PHONE.sub("[PHONE]", text)


def truncate(text: str) -> Tuple[str, bool]:
    """Cat o ranh gioi dong gan MAX_CHARS nhat."""
    if len(text) <= MAX_CHARS:
        return text, False
    cut = text.rfind("\n", 0, MAX_CHARS)
    return text[: cut if cut > 0 else MAX_CHARS], True


def has_injection(value: str) -> bool:
    return bool(_INJECTION.search(value))


def prepare_cv_text(data: bytes) -> Tuple[str, bool]:
    """Bytes -> (text da lam sach + redact, truncated). Loi neu qua ngan (nghi anh scan)."""
    text = normalize(extract_text(data))
    if len(text) < MIN_CHARS:
        raise ApiError(422, "cv_unreadable", f"Extracted text shorter than {MIN_CHARS} characters",
                       {"chars": len(text)})
    text, truncated = truncate(redact_pii(text))
    return text, truncated
